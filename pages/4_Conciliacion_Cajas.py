import streamlit as st
import pandas as pd
from datetime import datetime
import io
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --- CONFIGURACIÓN DE PÁGINA ---
st.set_page_config(
    page_title="Conciliación de Cajas - Portal Popsy",
    page_icon="🍦",
    layout="wide"
)

# --- ESTILOS CSS PERSONALIZADOS ---
st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Montserrat:wght@400;600;700;800&display=swap');
    
    html, body, [class*="css"] {
        font-family: 'Montserrat', sans-serif;
    }
    
    .top-bar {
        display: flex;
        justify-content: space-between;
        align-items: center;
        background-color: var(--background-secondary, #FFFFFF);
        padding: 12px 25px;
        border-bottom: 1px solid #E5E7EB;
        margin-bottom: 20px;
        border-radius: 12px;
        box-shadow: 0 2px 6px rgba(0,0,0,0.05);
    }
    .brand-logo {
        background-color: #FFF0F5;
        color: #C2185B;
        padding: 6px 18px;
        border-radius: 20px;
        font-weight: 800;
        font-size: 14px;
        letter-spacing: 0.5px;
    }
    .user-pill {
        display: flex;
        align-items: center;
        gap: 10px;
        background: var(--background-secondary, #FFFFFF);
        border: 1px solid #E5E7EB;
        padding: 6px 14px;
        border-radius: 30px;
        box-shadow: 0 2px 4px rgba(0,0,0,0.03);
    }
    .avatar {
        background-color: #C2185B;
        color: white;
        width: 32px;
        height: 32px;
        border-radius: 50%;
        display: flex;
        align-items: center;
        justify-content: center;
        font-weight: bold;
        font-size: 14px;
    }
    .user-info {
        display: flex;
        flex-direction: column;
        line-height: 1.2;
    }
    .user-name {
        font-weight: 700;
        font-size: 13px;
    }
    .user-role {
        font-size: 10px;
        font-weight: 700;
        color: #C2185B;
        letter-spacing: 0.5px;
    }
    .hero-banner {
        background: linear-gradient(135deg, #111827 0%, #1F2937 60%, #374151 100%);
        color: white;
        padding: 30px 35px;
        border-radius: 18px;
        margin-bottom: 25px;
        box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.2);
    }
    .badge-tag {
        background-color: rgba(194, 24, 91, 0.25);
        color: #F472B6;
        padding: 5px 14px;
        border-radius: 12px;
        font-size: 12px;
        font-weight: 700;
        display: inline-block;
        margin-bottom: 12px;
        border: 1px solid rgba(244, 114, 182, 0.3);
    }
    .hero-title {
        font-size: 26px;
        font-weight: 800;
        margin: 0 0 6px 0;
        color: #FFFFFF !important;
    }
    .hero-sub {
        color: #D1D5DB !important;
        font-size: 14px;
        margin: 0;
        font-style: italic;
    }
    .stButton>button {
        background-color: #C2185B !important;
        color: white !important;
        border-radius: 10px !important;
        font-weight: 700 !important;
        border: none !important;
        padding: 0.6rem 1.2rem !important;
    }
    .stButton>button:hover {
        background-color: #8E0038 !important;
    }
    </style>
""", unsafe_allow_html=True)

# --- VALIDACIÓN DE SESIÓN ---
if "autenticado" not in st.session_state or not st.session_state.autenticado:
    st.markdown("""
        <div style="background-color: #FFF3CD; padding: 25px; border-radius: 12px; border-left: 6px solid #FFEEBA; text-align: center; margin-top: 50px;">
            <h3 style="color: #856404; margin-top: 0;">⚠️ Acceso Restringido</h3>
            <p style="font-size: 1.1rem; color: #533F03;">Por favor, regresa a la página principal (<b>app.py</b>) e inicia sesión con tu correo corporativo para acceder a este módulo.</p>
        </div>
    """, unsafe_allow_html=True)
else:
    nombre_corto = st.session_state.get('usuario_nombre', st.session_state.correo_usuario.split('@')[0].replace('.', ' ').title())

    # --- TOP BAR SUPERIOR ---
    st.markdown(f"""
        <div class="top-bar">
            <div class="brand-logo">CONCILIACIÓN DE CAJAS</div>
            <div class="user-pill">
                <div class="avatar">{nombre_corto[0]}</div>
                <div class="user-info">
                    <span class="user-name">{nombre_corto}</span>
                    <span class="user-role">COLABORADOR POPSY</span>
                </div>
            </div>
        </div>
    """, unsafe_allow_html=True)

    # --- HERO BANNER ---
    fecha_actual = datetime.now().strftime('%A, %d de %B de %Y').capitalize()
    st.markdown(f"""
        <div class="hero-banner">
            <div class="badge-tag">🍦 CONCILIACIÓN DE CAJAS • {fecha_actual}</div>
            <h1 class="hero-title">¡Hola, {nombre_corto.split()[0]}! Módulo de Conciliación de Cajas</h1>
            <p class="hero-sub">"Hacemos del mundo un lugar más dulce, feliz y divertido"</p>
        </div>
    """, unsafe_allow_html=True)

    # --- SIDEBAR: SELECCIÓN DE EMPRESA Y CONFIGURACIÓN ---
    st.sidebar.markdown("### 🏢 Selección de Empresa")
    sociedad_seleccionada = st.sidebar.selectbox(
        "Selecciona la Sociedad:",
        [
            "Comercial Allan (CO01)",
            "Expertos en Café (CO02)",
            "Expertos en Café Peru (PE02 - En Proceso)",
            "Franquicias Unidad del Peru (PE01 - En Proceso)",
            "Panameña de Helados (PA01 - En Proceso)"
        ]
    )

    st.sidebar.markdown("---")
    st.sidebar.markdown("### ⚙️ Configuración")
    if st.sidebar.button("Cerrar Sesión", use_container_width=True):
        st.session_state.autenticado = False
        st.rerun()

    if "CO01" in sociedad_seleccionada:
        codigo_sociedad = "CO01"
    elif "CO02" in sociedad_seleccionada:
        codigo_sociedad = "CO02"
    else:
        codigo_sociedad = "GEN"

    def aplicar_estilos_excel_bytes(wb):
        NAVY_HEADER = "003366"
        WHITE_TEXT = "FFFFFF"
        ZEBRA_FILL = "F2F5F9"
        BORDER_COLOR = "D9D9D9"
        ACCENT_BLUE = "D9E1F2"
        
        thin_border = Border(
            left=Side(style='thin', color=BORDER_COLOR), right=Side(style='thin', color=BORDER_COLOR),
            top=Side(style='thin', color=BORDER_COLOR), bottom=Side(style='thin', color=BORDER_COLOR)
        )
        
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            ws.views.sheetView[0].showGridLines = True
            
            if 'Resumen' in sheetname:
                ws.column_dimensions['A'].width = 35
                ws.column_dimensions['B'].width = 25
                ws.merge_cells('A1:B1')
                title_cell = ws['A1']
                title_cell.value = "INFORME EJECUTIVO DE CONCILIACIÓN DE CAJAS"
                title_cell.font = Font(name='Segoe UI', size=13, bold=True, color=WHITE_TEXT)
                title_cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
                title_cell.alignment = Alignment(horizontal='center', vertical='center')
                ws.row_dimensions[1].height = 35
                
                for row in range(3, ws.max_row + 1):
                    cell_a, cell_b = ws[f'A{row}'], ws[f'B{row}']
                    cell_a.font = Font(name='Segoe UI', size=11, bold=(row == ws.max_row))
                    cell_b.font = Font(name='Segoe UI', size=11, bold=(row == ws.max_row))
                    cell_a.border, cell_b.border = thin_border, thin_border
                    cell_b.alignment = Alignment(horizontal='right')
                    if row % 2 == 0:
                        cell_a.fill = cell_b.fill = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
                    if row == ws.max_row:
                        cell_a.fill = cell_b.fill = PatternFill(start_color=ACCENT_BLUE, end_color=ACCENT_BLUE, fill_type="solid")
            else:
                ws.row_dimensions[1].height = 26
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row=1, column=col)
                    cell.font = Font(name='Segoe UI', size=10, bold=True, color=WHITE_TEXT)
                    cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.border = thin_border
                    
                for row in range(2, ws.max_row + 1):
                    ws.row_dimensions[row].height = 20
                    for col in range(1, ws.max_column + 1):
                        cell = ws.cell(row=row, column=col)
                        cell.font = Font(name='Segoe UI', size=10)
                        cell.border = thin_border
                        if row % 2 == 0:
                            cell.fill = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
                            
                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 40)
        return wb

    # --- INTERFAZ PRINCIPAL DE CONCILIACIÓN DE CAJAS ---
    st.markdown(f"### 💰 Conciliación y Arqueo de Cajas • {sociedad_seleccionada}")
    st.markdown("Carga el archivo de control de cajas (con columnas de **PDV**, **Concepto**, y **Fecha Control**) para realizar la validación y obtener el consolidado unificado por punto de venta.")
    
    archivo_cajas = st.file_uploader("Cargar Archivo de Cajas (Excel o CSV)", type=["xlsx", "csv"], key="cajas_page_uploader")
    
    if archivo_cajas:
        try:
            if archivo_cajas.name.endswith('.csv'):
                df_cajas = pd.read_csv(archivo_cajas)
            else:
                df_cajas = pd.read_excel(archivo_cajas)
            
            df_cajas.columns = [str(c).strip().upper() for c in df_cajas.columns]
            
            st.success("✅ Archivo de cajas cargado exitosamente.")
            with st.expander("🔍 Ver vista previa de los datos crudos"):
                st.dataframe(df_cajas.head(5), use_container_width=True)
            
            if st.button("🚀 Ejecutar Análisis y Consolidado por PDV"):
                with st.spinner("Procesando arqueos, unificando por PDV y calculando diferencias..."):
                    
                    # Detección inteligente de columnas clave
                    col_pdv = next((c for c in df_cajas.columns if 'PDV' in c or 'PUNTO' in c or 'TIENDA' in c or 'CENTRO' in c), df_cajas.columns[0])
                    col_val_sistema = next((c for c in df_cajas.columns if 'SISTEMA' in c or 'TEORICO' in c or 'BASE' in c or 'VALOR' in c), df_cajas.columns[min(2, len(df_cajas.columns)-1)])
                    col_val_fisico = next((c for c in df_cajas.columns if 'FISICO' in c or 'REAL' in c or 'CONTEO' in c or 'CONTADO' in c), df_cajas.columns[min(3, len(df_cajas.columns)-1)])
                    
                    # Limpiar y convertir valores numéricos
                    df_cajas['VAL_SISTEMA'] = pd.to_numeric(df_cajas[col_val_sistema].astype(str).str.replace(r'[^0-9.-]', '', regex=True), errors='coerce').fillna(0)
                    df_cajas['VAL_FISICO'] = pd.to_numeric(df_cajas[col_val_fisico].astype(str).str.replace(r'[^0-9.-]', '', regex=True), errors='coerce').fillna(0)
                    
                    # Cálculo de diferencia individual
                    df_cajas['DIFERENCIA'] = df_cajas['VAL_FISICO'] - df_cajas['VAL_SISTEMA']
                    
                    def clasificar_estado(row):
                        dif = row['DIFERENCIA']
                        if abs(dif) < 1:
                            return "CUADRADO (OK)"
                        elif dif > 0:
                            return "SOBRANTE"
                        else:
                            return "FALTANTE"
                            
                    df_cajas['ESTADO_ARQUEO'] = df_cajas.apply(clasificar_estado, axis=1)
                    
                    # --- UNIFICACIÓN / CONSOLIDADO POR PDV ---
                    df_pdv_resumen = df_cajas.groupby(col_pdv).agg(
                        Total_Registros=('VAL_SISTEMA', 'count'),
                        Total_Sistema=('VAL_SISTEMA', 'sum'),
                        Total_Fisico=('VAL_FISICO', 'sum'),
                        Diferencia_Neta=('DIFERENCIA', 'sum')
                    ).reset_index()
                    
                    def estado_pdv(row):
                        dif = row['Diferencia_Neta']
                        if abs(dif) < 1:
                            return "CUADRADO (OK)"
                        elif dif > 0:
                            return "SOBRANTE NETO"
                        else:
                            return "FALTANTE NETO"
                            
                    df_pdv_resumen['ESTADO_PDV'] = df_pdv_resumen.apply(estado_pdv, axis=1)
                    
                    total_pdvs = len(df_pdv_resumen)
                    monto_neto_global = df_cajas['DIFERENCIA'].sum()
                    
                    st.markdown("### 📊 Resultados Generales")
                    mc1, mc2, mc3 = st.columns(3)
                    mc1.metric(label="Total PDV Analizados", value=f"{total_pdvs} tiendas")
                    mc2.metric(label="Total Registros Procesados", value=f"{len(df_cajas)} partidas")
                    mc3.metric(label="Diferencia Neta Global", value=f"${monto_neto_global:,.2f}", delta="Balance general", delta_color="off")
                    
                    # Pestañas para ver el Consolidado por PDV y el Detalle Completo
                    tab_pdv, tab_detalle = st.tabs(["🏢 Consolidado Unificado por PDV", "📋 Detalle Completo de Cajas"])
                    
                    with tab_pdv:
                        st.markdown(f"#### Resumen Unificado por {col_pdv}")
                        st.dataframe(df_pdv_resumen, use_container_width=True)
                        
                    with tab_detalle:
                        st.markdown("#### Detalle de Arqueos, Conceptos y Fecha Control")
                        st.dataframe(df_cajas, use_container_width=True)
                    
                    # Exportar informe en Excel con múltiples pestañas estilizadas
                    output_cajas = io.BytesIO()
                    with pd.ExcelWriter(output_cajas, engine='openpyxl') as writer:
                        pd.DataFrame({
                            "Métrica Caja": ["Sociedad", "Fecha Proceso", "Total PDV Evaluados", "Total Partidas", "Diferencia Neta Total"],
                            "Valor": [codigo_sociedad, datetime.now().strftime('%Y-%m-%d %H:%M'), total_pdvs, len(df_cajas), monto_neto_global]
                        }).to_excel(writer, sheet_name='Resumen_Cajas', index=False, startrow=1)
                        
                        df_pdv_resumen.to_excel(writer, sheet_name='Consolidado_Por_PDV', index=False)
                        df_cajas.to_excel(writer, sheet_name='Detalle_Completo', index=False)
                        
                    output_cajas.seek(0)
                    wb_c = openpyxl.load_workbook(output_cajas)
                    wb_c_estilizado = aplicar_estilos_excel_bytes(wb_c)
                    
                    final_output_cajas = io.BytesIO()
                    wb_c_estilizado.save(final_output_cajas)
                    final_output_cajas.seek(0)
                    
                    st.markdown("---")
                    st.download_button(
                        label="📥 Descargar Informe Completo y Consolidado por PDV (.xlsx)",
                        data=final_output_cajas,
                        file_name=f"Informe_Consolidado_Cajas_{codigo_sociedad}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
                    
        except Exception as e:
            st.error(f"Error procesando el archivo de cajas: {e}")
    else:
        st.info("📌 Por favor carga el archivo de control de cajas para generar el unificado por punto de venta.")
